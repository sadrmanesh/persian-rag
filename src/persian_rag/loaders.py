"""PDF and DOCX loaders with Persian-aware cleanup.

Text extracted from PDFs often needs repair before it can be indexed:

* Arabic presentation forms (contextual glyphs such as U+FEE3) instead of
  ordinary letters.
* Lines stored in visual order, so every Persian word appears reversed.
* Bidirectional control marks, soft hyphens, and other invisible characters.

:func:`clean_extracted_text` handles all three. Visual-order lines are
detected per page, reversed, and the left-to-right runs inside them (numbers
and Latin words) are restored. Final letter normalization (Arabic Yeh and
Kaf, digits, spacing) is left to :func:`persian_rag.normalize.normalize`,
which runs during ingestion.

Install the extras with ``pip install 'persian-rag[pdf]'`` (pypdf) or
``pip install 'persian-rag[docx]'`` (python-docx).
"""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Callable, Iterator, Sequence
from pathlib import Path
from typing import Any, Literal

from persian_rag.documents import SourceDocument
from persian_rag.errors import MissingDependencyError, PersianRagError

RtlFix = Literal["auto", "reverse", "none"]

# Bidi embedding, override, and isolate controls, plus marks and soft hyphen.
_INVISIBLE_RE = re.compile("[\u00ad\u061c\u200e\u200f\u202a-\u202e\u2066-\u2069\ufeff]")
_CONTROL_RE = re.compile("[\u0000-\u0008\u000b\u000c\u000e-\u001f\u007f]")
_LTR_RUN_RE = re.compile(
    r"[A-Za-z0-9\u06f0-\u06f9\u0660-\u0669]+(?:[.,:/%@_\-][A-Za-z0-9\u06f0-\u06f9\u0660-\u0669]+)*"
)
_ARABIC_LETTER_RE = re.compile("[\u0600-\u06ff\ufb50-\ufdff\ufe70-\ufeff]")
_WORD_RE = re.compile("[\u0600-\u06ff\ufb50-\ufdff\ufe70-\ufeff]+")
_MIRROR = str.maketrans("()[]{}<>«»", ")(][}{><»«")

# Frequent short Persian words and their reversed spellings.
_COMMON_WORDS = frozenset({"از", "به", "که", "در", "این", "را", "با", "است", "برای", "یک", "آن"})
_REVERSED_COMMON = frozenset(word[::-1] for word in _COMMON_WORDS) - _COMMON_WORDS


def _form(char: str) -> str | None:
    code = ord(char)
    if not (0xFB50 <= code <= 0xFDFF or 0xFE70 <= code <= 0xFEFF):
        return None
    name = unicodedata.name(char, "")
    for form in ("INITIAL", "MEDIAL", "FINAL", "ISOLATED"):
        if f"{form} FORM" in name:
            return form
    return None


def visual_order_score(text: str) -> int:
    """Evidence that ``text`` is stored in visual (reversed) order.

    Positive values suggest visual order, negative values logical order. The
    score combines two signals: which presentation forms open and close each
    word, and whether frequent Persian words appear forwards or reversed.
    """

    score = 0
    for word in _WORD_RE.findall(text):
        if len(word) > 1:
            first, last = _form(word[0]), _form(word[-1])
            if first == "FINAL" or last == "INITIAL":
                score += 1
            if first == "INITIAL" or last == "FINAL":
                score -= 1
        folded = unicodedata.normalize("NFKC", word)
        if folded in _REVERSED_COMMON:
            score += 2
        elif folded in _COMMON_WORDS:
            score -= 2
    return score


def looks_visual_order(text: str) -> bool:
    """Return True when ``text`` is probably stored in visual order."""

    return visual_order_score(text) > 0


def reverse_rtl_line(line: str) -> str:
    """Convert one visual-order line to logical order.

    The characters are reversed, brackets are mirrored, and runs of digits or
    Latin letters are flipped back so ``"2026"`` stays ``"2026"``.
    """

    flipped = line[::-1].translate(_MIRROR)
    return _LTR_RUN_RE.sub(lambda match: match.group(0)[::-1], flipped)


def fold_presentation_forms(text: str) -> str:
    """Replace Arabic presentation forms with ordinary letters (NFKC per glyph)."""

    return "".join(
        unicodedata.normalize("NFKC", char) if _form(char) is not None else char for char in text
    )


def clean_extracted_text(text: str, *, rtl_fix: RtlFix = "auto") -> str:
    """Clean text extracted from a PDF or DOCX page.

    ``rtl_fix`` controls visual-order repair: ``"auto"`` detects it per
    page, ``"reverse"`` always reverses lines containing Persian letters,
    and ``"none"`` leaves the order unchanged.
    """

    if rtl_fix not in ("auto", "reverse", "none"):
        raise ValueError("rtl_fix must be 'auto', 'reverse', or 'none'")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _INVISIBLE_RE.sub("", text)
    text = _CONTROL_RE.sub(" ", text)
    reverse = rtl_fix == "reverse" or (rtl_fix == "auto" and looks_visual_order(text))
    lines = text.split("\n")
    if reverse:
        lines = [
            reverse_rtl_line(line) if _ARABIC_LETTER_RE.search(line) else line for line in lines
        ]
    cleaned = "\n".join(line.rstrip() for line in lines)
    cleaned = fold_presentation_forms(cleaned)
    return re.sub(r"\n{3,}", "\n\n", cleaned).strip()


PdfReaderFactory = Callable[[Path, str | None], Any]
DocxFactory = Callable[[Path], Any]


def _default_pdf_reader(path: Path, password: str | None) -> Any:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingDependencyError(
            "pypdf is not installed. Install it with: pip install 'persian-rag[pdf]'"
        ) from exc
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        if password is None:
            raise PersianRagError(f"{path} is encrypted; pass a password")
        reader.decrypt(password)
    return reader


def _pdf_pages(path: Path, password: str | None, factory: PdfReaderFactory | None) -> list[str]:
    opener = factory or _default_pdf_reader
    try:
        reader = opener(path, password)
        return [str(page.extract_text() or "") for page in reader.pages]
    except (PersianRagError, ImportError):
        raise
    except Exception as exc:
        raise PersianRagError(f"Could not read PDF {path}: {exc}") from exc


def load_pdf(
    path: str | Path,
    *,
    document_id: str | None = None,
    rtl_fix: RtlFix = "auto",
    split_pages: bool = False,
    password: str | None = None,
    reader_factory: PdfReaderFactory | None = None,
) -> list[SourceDocument]:
    """Extract and clean the text of a PDF.

    Returns one document for the whole file, or one per non-empty page when
    ``split_pages`` is true (ids end with ``#page=N``). ``reader_factory`` lets
    callers plug in another extractor; it receives the path and password and
    returns an object with a ``pages`` list whose items have
    ``extract_text()``.
    """

    source = Path(path)
    doc_id = document_id or source.name
    pages = [
        clean_extracted_text(page, rtl_fix=rtl_fix)
        for page in _pdf_pages(source, password, reader_factory)
    ]
    base = {"source": doc_id, "filename": source.name, "format": "pdf"}
    if split_pages:
        return [
            SourceDocument(
                id=f"{doc_id}#page={number}",
                text=text,
                metadata={**base, "page": number, "page_count": len(pages)},
            )
            for number, text in enumerate(pages, start=1)
            if text
        ]
    text = "\n\n".join(page for page in pages if page)
    if not text:
        return []
    return [SourceDocument(id=doc_id, text=text, metadata={**base, "page_count": len(pages)})]


def _default_docx(path: Path) -> Any:
    try:
        import docx
    except ImportError as exc:
        raise MissingDependencyError(
            "python-docx is not installed. Install it with: pip install 'persian-rag[docx]'"
        ) from exc
    return docx.Document(str(path))


def _docx_blocks(document: Any) -> Iterator[str]:
    iterate = getattr(document, "iter_inner_content", None)
    items: Sequence[Any] = (
        list(iterate())
        if callable(iterate)
        else [
            *document.paragraphs,
            *document.tables,
        ]
    )
    for item in items:
        if hasattr(item, "rows"):
            for row in item.rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = " ".join(cell.text.split())
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
                if cells:
                    yield " | ".join(cells)
        else:
            yield str(item.text)


def load_docx(
    path: str | Path,
    *,
    document_id: str | None = None,
    rtl_fix: RtlFix = "none",
    document_factory: DocxFactory | None = None,
) -> list[SourceDocument]:
    """Extract paragraphs and tables from a Word document in body order.

    Word stores text in logical order, so ``rtl_fix`` defaults to ``"none"``.
    Table rows become ``cell | cell`` lines; merged cells are not repeated.
    """

    source = Path(path)
    doc_id = document_id or source.name
    opener = document_factory or _default_docx
    try:
        document = opener(source)
        blocks = list(_docx_blocks(document))
    except (PersianRagError, ImportError):
        raise
    except Exception as exc:
        raise PersianRagError(f"Could not read DOCX {source}: {exc}") from exc
    text = clean_extracted_text("\n".join(blocks), rtl_fix=rtl_fix)
    if not text:
        return []
    return [
        SourceDocument(
            id=doc_id,
            text=text,
            metadata={"source": doc_id, "filename": source.name, "format": "docx"},
        )
    ]
